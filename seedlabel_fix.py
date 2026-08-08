#!/usr/bin/env python3
"""
seedlabel_fix.py -- remove the causal attribution to the seed. Nothing else.

ESTABLISHED FACT THIS RESTS ON. env_6g_v5.py:132 builds the channel RNG as
np.random.default_rng(seed). Both trainers construct the environment WITHOUT a seed --
run_multiseed_rel.py:43 and train_branched_nyu.py:39 -- so every training episode draws its
shadow fading, its small-scale fading and its trajectory start from fresh OS entropy. set_seed(s)
controls the epsilon-greedy draw, the numpy legacy global, torch init, and which trajectory is
sampled. It does NOT control the channel the agent trains against. Confirmed empirically: two
trainings of the guarded arm at seed 0, identical config and budget, produce different weights.

SCOPE DISCIPLINE. A separate experiment (5 trainings at the SAME seed) is running and will say
HOW MUCH of the dispersion the seed accounts for. This edit deliberately does NOT anticipate it.
Every replacement below is true whether the seed turns out to explain most of the spread, some,
or none -- so nothing here has to be rewritten when that number lands. The affirmative reframing
("a retraining lottery, not a seed lottery"), the phrase "reliability seed lottery", and the
title are all LEFT ALONE for that reason.

WHAT CHANGES
  1. Five causal statements that are false as written today.
  2. "training-seed component" -> "between-training component" throughout. This is what the
     quantity has always been mathematically: sigma_train is estimated across independent
     TRAINING RUNS, and the decomposition never cared which entropy source drove it.
  3. One disclosure sentence in Section VIII where sigma_train is defined, and one in the
     reproducibility section naming where the unseededness lives.

WHAT IS DELIBERATELY NOT CHANGED
  * "training seeds" (plural) describing the DESIGN -- "5 independent training seeds per arm" is
    an accurate description of what was run and must survive. The regexes below are written to
    match the hyphenated compound and the singular causal phrase, never the plural noun.
  * The [219] claim "training is not venue-portable". On a full reading the manuscript justifies
    it by the unseeded protocol, not by a single cross-venue training pair -- and that
    justification is exactly what has now been confirmed. It is sharpened, not withdrawn.
"""
import hashlib
import re
import sys

ORIG = '/home/claude/work/ms/DnDQN_manuscript_updated_IXrenumbered.docx'
ORIG_MD5 = 'ddea24ec8ad709ab02356486b18c46e2'
WORK = '/home/claude/work/ms_edit/DnDQN_manuscript_WORKING.docx'
WORK_IN = '325ac4d7d7c5d0c7ff881a2a83b70a63'


def md5(p):
    h = hashlib.md5()
    with open(p, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


from docx import Document                                            # noqa: E402
from docx.oxml.ns import qn                                          # noqa: E402
from docx.table import Table                                         # noqa: E402
from docx.text.paragraph import Paragraph                            # noqa: E402


def blocks(doc):
    out = []
    for ch in doc.element.body.iterchildren():
        if ch.tag == qn('w:p'):
            out.append(Paragraph(ch, doc))
        elif ch.tag == qn('w:tbl'):
            out.append(Table(ch, doc))
    return out


def all_paras(B):
    """Every paragraph in the document, including those inside table cells."""
    for i, b in enumerate(B):
        if isinstance(b, Table):
            for row in b.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        yield i, p
        else:
            yield i, b


def tol(s):
    out = []
    for ch in s:
        if ch == "'":
            out.append("['’‘]")
        elif ch == '"':
            out.append('["“”]')
        elif ch == '-':
            out.append('[-‐‑–—]')
        elif ch == ' ':
            out.append(r'\s+')
        else:
            out.append(re.escape(ch))
    return ''.join(out)


def replace_span(p, rx, repl):
    runs = list(p.runs)
    if not runs:
        return False
    text = ''.join(r.text for r in runs)
    m = rx.search(text)
    if not m:
        return False
    s, e = m.span()
    if callable(repl):
        repl = repl(m)
    pos, first, edits = 0, None, []
    for r in runs:
        L = len(r.text)
        a, b = pos, pos + L
        lo, hi = max(a, s), min(b, e)
        if lo < hi:
            head, tail = r.text[:lo - a], r.text[hi - a:]
            if first is None:
                first = r
                edits.append((r, head + repl + tail))
            else:
                edits.append((r, head + tail))
        pos = b
    for r, t in edits:
        r.text = t
    return True


TARGETED = [
 (5, 'abstract: causal attribution',
  'a 3.3-fold spread produced by nothing but the seed',
  'a 3.3-fold spread produced by retraining alone'),
 (5, 'abstract: causal attribution',
  'decided by the random seed rather than by the design',
  'decided by the training run rather than by the design'),
 (5, 'abstract: seed-to-seed',
  'so seed-to-seed variation is 22.6 times larger',
  'so run-to-run variation is 22.6 times larger'),
 (94, 'Section VI: "only the seed changes"',
  'are all the same; only the random seed changes',
  'are all the same; the trainings differ only in their random draws'),
 (97, 'Section VI: causal attribution',
  'a 3.3× spread produced by nothing but the seed',
  'a 3.3× spread produced by retraining alone'),
 (156, 'IX-B: causal attribution',
  'the reliability variation attributable to the training seed is below',
  'the reliability variation attributable to retraining is below'),
 (132, 'Section VIII: define what sigma_train spans',
  'and is the larger quantity.',
  'and is the larger quantity. Independent retrainings differ in their declared seed and in '
  'their training-time channel draws alike — the environment is constructed without a '
  'channel seed during training — so σ_train is the dispersion across independent '
  'training runs, not the effect of the seed variable in isolation.'),
 (219, 'reproducibility: name where the unseededness lives',
  'whereas the unseeded protocol cannot, which is the reason the seeded one was adopted.',
  'whereas the unseeded protocol cannot, which is the reason the seeded one was adopted. The '
  'unseededness sits in training: both training drivers construct the environment without a '
  'channel seed, so each training episode draws its fading and its trajectory start from fresh '
  'entropy. Independent trainings therefore differ in more than their declared seed, which is '
  'why every dispersion in this paper is reported as a between-training component rather than '
  'as an effect of the seed variable.'),
]

# global renames, applied over every paragraph including table cells
GLOBAL = [
 (r'training[-‐‑‒–—]seed\s+(variance\s+)?component', lambda m: 'between-training %scomponent'
  % (m.group(1) or ''), 'training-seed component -> between-training component'),
 (r'training[-‐‑‒–—]seed\s+variance\b', lambda m: 'between-training variance',
  'training-seed variance -> between-training variance'),
 (r'training[-‐‑‒–—]seed\s+uncertainty', lambda m: 'between-training uncertainty',
  'training-seed uncertainty -> between-training uncertainty'),
]

NUM = re.compile(r'\d+(?:\.\d+)?')


def numerals(doc):
    """Every numeral in the document, from the raw <w:t> nodes in document order.

    The first version walked paragraphs and de-duplicated merged table cells by id(cell._tc).
    lxml creates element proxies on demand and frees them, so id() is not stable between calls:
    two traversals of an UNCHANGED document returned different multisets, and the invariance
    gate reported 11 numerals lost and 20 gained on edits that touched none of them. The text
    nodes themselves are canonical and ordered; use those."""
    acc = []
    for t in doc.element.body.iter(qn('w:t')):
        if t.text:
            acc += NUM.findall(t.text)
    return sorted(acc)


def main():
    bar = '=' * 104
    print(bar)
    print('SEED-CAUSALITY CORRECTION -- the seed is a label, not the sole cause')
    print(bar)
    for nm, path, want in (('original', ORIG, ORIG_MD5), ('working IN', WORK, WORK_IN)):
        got = md5(path)
        print('  ' + nm.ljust(12) + got + '  ' + ('OK' if got == want else 'MISMATCH -- STOP'))
        if got != want:
            return 2

    doc = Document(WORK)
    B = blocks(doc)
    before = numerals(doc)

    print('\n  PASS 0 -- gates')
    print('  ' + '-' * 100)
    bad = 0
    plan = []
    for bi, tag, old, new in TARGETED:
        if bi >= len(B) or isinstance(B[bi], Table):
            print('    [%3d] NOT A PARAGRAPH  %s' % (bi, tag))
            bad += 1
            continue
        rx = re.compile(tol(old))
        n = len(rx.findall(B[bi].text))
        ok = n == 1
        bad += 0 if ok else 1
        print('    [%3d] %s  %-46s %s...' % (bi, 'OK ' if ok else '*** %d hits ***' % n,
                                             tag, old[:40]))
        if ok:
            plan.append((bi, rx, new))

    counts = {}
    for pat, _, label in GLOBAL:
        rx = re.compile(pat)
        c = sum(len(rx.findall(p.text)) for _, p in all_paras(B))
        counts[pat] = c
        print('    glob  %s  %-52s %d occurrence(s)' % ('OK ' if c else '*** 0 ***', label, c))
        bad += 0 if c else 1

    # protect the plural design noun
    plural = sum(len(re.findall(r'training seeds', p.text, re.I)) for _, p in all_paras(B))
    print('    guard "training seeds" (plural, DESIGN — must survive): %d occurrence(s)' % plural)
    if bad:
        print('\n  %d gate failure(s). NOTHING WRITTEN.' % bad)
        return 3

    print('\n  PASS 1 -- applying')
    for bi, rx, new in plan:
        if not replace_span(B[bi], rx, new):
            print('    [%d] APPLY FAILED' % bi)
            return 4
    for pat, repl, label in GLOBAL:
        rx = re.compile(pat)
        done = 0
        for _, p in all_paras(B):
            for _ in range(12):
                if not rx.search(''.join(r.text for r in p.runs)):
                    break
                replace_span(p, rx, repl)
                done += 1
        print('    %-54s %d replacement(s)' % (label, done))
    print('    %d targeted edit(s)' % len(plan))

    print('\n  PASS 2 -- post-conditions')
    left = [(bi, p.text) for bi, p in all_paras(B) if re.search(r'training[-\u2010\u2011\u2012\u2013\u2014]seed', p.text)]
    print('    "training-seed" remaining : %d %s'
          % (len(left), 'OK' if not left else '*** ' + str([b for b, _ in left]) + ' ***'))
    plural2 = sum(len(re.findall(r'training seeds', p.text, re.I)) for _, p in all_paras(B))
    print('    "training seeds" (plural) : %d -> %d  %s'
          % (plural, plural2, 'PRESERVED' if plural == plural2 else '*** DAMAGED ***'))
    after = numerals(doc)
    same = before == after
    print('    numerals %d -> %d  %s' % (len(before), len(after),
                                         'IDENTICAL' if same else '*** CHANGED ***'))
    if not same:
        from collections import Counter
        c1, c2 = Counter(before), Counter(after)
        print('      LOST  : %s' % sorted((c1 - c2).elements()))
        print('      GAINED: %s' % sorted((c2 - c1).elements()))
    if left or plural != plural2 or not same:
        print('    NOTHING WRITTEN.')
        return 5

    doc.save(WORK)
    print('\n  original re-verified : %s  %s'
          % (md5(ORIG), 'UNCHANGED' if md5(ORIG) == ORIG_MD5 else '*** WRITTEN TO ***'))
    print('  working md5 OUT      : %s' % md5(WORK))

    B2 = blocks(Document(WORK))
    print('\n  READ BACK')
    for bi in (5, 94, 97, 132, 156, 219):
        t = B2[bi].text
        print('\n  [%d] %s' % (bi, t[:600] + ('...' if len(t) > 600 else '')))
    print('\n' + bar)
    return 0


if __name__ == '__main__':
    sys.exit(main())
